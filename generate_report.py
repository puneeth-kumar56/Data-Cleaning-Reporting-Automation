"""
generate_report.py
-------------------
Builds all charts (PNG) and the final Excel workbook
(reports/data_cleaning_report.xlsx) from CleaningResult and AnalysisResult.
"""

import textwrap
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")          # non-interactive backend (no display needed)
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
import openpyxl
import pandas as pd
import seaborn as sns
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import (Alignment, Border, Font, GradientFill,
                              PatternFill, Side)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT        = Path(__file__).resolve().parents[1]
CHARTS_DIR  = ROOT / "reports" / "charts"
REPORT_FILE = ROOT / "reports" / "data_cleaning_report.xlsx"
SUMMARY_TXT = ROOT / "reports" / "data_quality_summary.txt"

# Matplotlib style
plt.rcParams.update({
    "figure.facecolor": "white",
    "axes.facecolor":   "white",
    "axes.grid":        True,
    "grid.alpha":       0.35,
    "font.family":      "DejaVu Sans",
})
PALETTE = sns.color_palette("Blues_d", 8)


# ══════════════════════════════════════════════════════════════════════════════
#  CHARTS
# ══════════════════════════════════════════════════════════════════════════════

def _save_fig(name: str) -> Path:
    CHARTS_DIR.mkdir(parents=True, exist_ok=True)
    path = CHARTS_DIR / name
    plt.savefig(path, dpi=120, bbox_inches="tight")
    plt.close()
    print(f"    Chart saved → {path.name}")
    return path


def chart_department_distribution(ana) -> Path:
    """Bar chart – student count per department."""
    fig, ax = plt.subplots(figsize=(8, 5))
    df = ana.dept_counts.sort_values("Student_Count", ascending=True)
    bars = ax.barh(df["Department"], df["Student_Count"],
                   color=sns.color_palette("Blues_d", len(df)))
    ax.bar_label(bars, padding=4, fontsize=10)
    ax.set_xlabel("Number of Students", fontsize=11)
    ax.set_ylabel("Department", fontsize=11)
    ax.set_title("Student Distribution by Department", fontsize=13, fontweight="bold")
    ax.set_xlim(0, df["Student_Count"].max() * 1.18)
    plt.tight_layout()
    return _save_fig("department_distribution.png")


def chart_avg_marks_by_dept(ana) -> Path:
    """Grouped bar chart – average marks per department."""
    df = ana.dept_avg_marks.set_index("Department")
    cols = ["Assignment_Marks", "Internal_Marks", "Final_Exam_Marks", "Total_Marks"]
    df = df[cols]

    fig, ax = plt.subplots(figsize=(10, 6))
    x     = np.arange(len(df))
    width = 0.2
    colors = ["#2196F3", "#4CAF50", "#FF9800", "#9C27B0"]

    for i, (col, color) in enumerate(zip(cols, colors)):
        rects = ax.bar(x + i * width, df[col], width, label=col.replace("_", " "),
                       color=color, alpha=0.82)
        ax.bar_label(rects, fmt="%.1f", padding=2, fontsize=7, rotation=45)

    ax.set_xlabel("Department", fontsize=11)
    ax.set_ylabel("Average Marks", fontsize=11)
    ax.set_title("Average Marks by Department", fontsize=13, fontweight="bold")
    ax.set_xticks(x + width * 1.5)
    ax.set_xticklabels(df.index, fontsize=10)
    ax.legend(fontsize=9, loc="upper right")
    plt.tight_layout()
    return _save_fig("average_marks_by_department.png")


def chart_grade_distribution(ana) -> Path:
    """Pie + bar combo – grade distribution."""
    df   = ana.grade_dist.sort_values("Grade")
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))

    # Pie
    wedge_props = {"linewidth": 1.2, "edgecolor": "white"}
    ax1.pie(df["Count"], labels=df["Grade"], autopct="%1.1f%%",
            colors=sns.color_palette("Set2", len(df)),
            wedgeprops=wedge_props, startangle=140)
    ax1.set_title("Grade Distribution (Pie)", fontsize=12, fontweight="bold")

    # Bar
    bars = ax2.bar(df["Grade"], df["Count"],
                   color=sns.color_palette("Set2", len(df)))
    ax2.bar_label(bars, padding=3, fontsize=10)
    ax2.set_xlabel("Grade", fontsize=11)
    ax2.set_ylabel("Count", fontsize=11)
    ax2.set_title("Grade Distribution (Bar)", fontsize=12, fontweight="bold")

    plt.suptitle("Student Grade Distribution", fontsize=14, fontweight="bold", y=1.01)
    plt.tight_layout()
    return _save_fig("grade_distribution.png")


def chart_missing_values_comparison(clean_res) -> Path:
    """Grouped bar – missing values before vs after cleaning."""
    labels  = ["Missing Before", "Missing After"]
    values  = [clean_res.missing_before, clean_res.missing_after]
    colors  = ["#EF5350", "#66BB6A"]

    fig, ax = plt.subplots(figsize=(6, 5))
    bars = ax.bar(labels, values, color=colors, width=0.4, edgecolor="white",
                  linewidth=1.2)
    ax.bar_label(bars, padding=4, fontsize=12, fontweight="bold")
    ax.set_ylabel("Number of Missing Values", fontsize=11)
    ax.set_title("Missing Values: Before vs After Cleaning", fontsize=12,
                 fontweight="bold")
    ax.set_ylim(0, max(values) * 1.3)
    plt.tight_layout()
    return _save_fig("missing_values_comparison.png")


def chart_attendance_distribution(ana) -> Path:
    """Histogram of attendance percentages."""
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.hist(ana.df["Attendance_Percentage"].dropna(), bins=20,
            color="#42A5F5", edgecolor="white", linewidth=0.8)
    ax.axvline(75, color="red", linestyle="--", linewidth=1.8, label="75% threshold")
    ax.set_xlabel("Attendance Percentage (%)", fontsize=11)
    ax.set_ylabel("Number of Students", fontsize=11)
    ax.set_title("Attendance Distribution", fontsize=13, fontweight="bold")
    ax.legend(fontsize=10)
    plt.tight_layout()
    return _save_fig("attendance_distribution.png")


def chart_dept_marks_comparison(ana) -> Path:
    """Line chart – total marks comparison across departments."""
    df = ana.dept_avg_marks.set_index("Department")
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(df.index, df["Total_Marks"], marker="o", color="#1565C0",
            linewidth=2.2, markersize=8, label="Avg Total Marks")
    for x_val, y_val in zip(df.index, df["Total_Marks"]):
        ax.annotate(f"{y_val:.1f}", (x_val, y_val),
                    textcoords="offset points", xytext=(0, 8),
                    ha="center", fontsize=9)
    ax.set_xlabel("Department", fontsize=11)
    ax.set_ylabel("Average Total Marks", fontsize=11)
    ax.set_title("Average Total Marks Comparison Across Departments",
                 fontsize=12, fontweight="bold")
    ax.legend(fontsize=10)
    ax.set_ylim(0, 100)
    plt.tight_layout()
    return _save_fig("dept_marks_comparison.png")


def generate_all_charts(clean_res, ana) -> dict[str, Path]:
    print("\n[Charts] Generating charts …")
    paths = {
        "dept_dist":     chart_department_distribution(ana),
        "avg_marks":     chart_avg_marks_by_dept(ana),
        "grade_dist":    chart_grade_distribution(ana),
        "missing_vals":  chart_missing_values_comparison(clean_res),
        "attendance":    chart_attendance_distribution(ana),
        "dept_compare":  chart_dept_marks_comparison(ana),
    }
    print(f"    {len(paths)} charts saved to {CHARTS_DIR}")
    return paths


# ══════════════════════════════════════════════════════════════════════════════
#  EXCEL WORKBOOK HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def _H(text: str, bold=True, size=12, color="FFFFFF", bg="1F4E79") -> dict:
    """Return a style dict for a header cell."""
    return {"value": text, "bold": bold, "size": size, "color": color, "bg": bg}

def _apply_cell(ws, row: int, col: int, value, bold=False, size=11,
                color="000000", bg=None, align="left", number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, size=size, color=color)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               wrap_text=True)
    if number_format:
        cell.number_format = number_format
    return cell


def _write_df_to_sheet(ws, df: pd.DataFrame, start_row=1,
                       header_bg="1F4E79", alt_bg="D9E1F2"):
    """Write a DataFrame to a worksheet with styled headers."""
    for col_idx, col_name in enumerate(df.columns, 1):
        _apply_cell(ws, start_row, col_idx, col_name,
                    bold=True, color="FFFFFF", bg=header_bg, align="center")

    for r_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
        fill_bg = alt_bg if r_idx % 2 == 0 else None
        for c_idx, value in enumerate(row, 1):
            v = None if (isinstance(value, float) and np.isnan(value)) else value
            _apply_cell(ws, r_idx, c_idx, v, bg=fill_bg)

    # Auto-width
    for col_cells in ws.iter_cols(min_row=start_row,
                                   max_row=start_row + len(df),
                                   max_col=len(df.columns)):
        max_len = max(
            len(str(cell.value or "")) for cell in col_cells
        ) + 4
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len, 35)

    # Filter + freeze
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes    = ws.cell(row=start_row + 1, column=1).coordinate


# ══════════════════════════════════════════════════════════════════════════════
#  SHEET BUILDERS
# ══════════════════════════════════════════════════════════════════════════════

def build_executive_summary(ws, clean_res, ana, now: str):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 24

    rows = [
        ("DATA CLEANING & REPORTING AUTOMATION", True, 16, "FFFFFF", "1F4E79"),
        ("Student Academic Performance Analysis", True, 12, "FFFFFF", "2E75B6"),
        (f"Report Generated: {now}", False, 10, "FFFFFF", "4472C4"),
        ("", False, 10, "000000", None),
        ("── RECORD STATISTICS ──────────────────────", True, 10, "000000", "BDD7EE"),
        ("Original Record Count",          False, 10, "000000", None),
        ("Final (Cleaned) Record Count",   False, 10, "000000", None),
        ("Duplicate Rows Removed",         False, 10, "000000", None),
        ("", False, 10, "000000", None),
        ("── MISSING VALUE STATISTICS ───────────────", True, 10, "000000", "BDD7EE"),
        ("Missing Values (Before Cleaning)", False, 10, "000000", None),
        ("Missing Values (After Cleaning)",  False, 10, "000000", None),
        ("", False, 10, "000000", None),
        ("── DATA QUALITY ────────────────────────────", True, 10, "000000", "BDD7EE"),
        ("Invalid Values Flagged",          False, 10, "000000", None),
        ("Categorical Values Corrected",    False, 10, "000000", None),
        ("Invalid Emails Flagged",          False, 10, "000000", None),
        ("", False, 10, "000000", None),
        ("── ACADEMIC PERFORMANCE ────────────────────", True, 10, "000000", "BDD7EE"),
        ("Total Students",                  False, 10, "000000", None),
        ("Average Age (years)",             False, 10, "000000", None),
        ("Average Attendance (%)",          False, 10, "000000", None),
        ("Average Total Marks",             False, 10, "000000", None),
        ("Highest Total Marks",             False, 10, "000000", None),
        ("Lowest Total Marks",              False, 10, "000000", None),
        ("Students with Attendance < 75%",  False, 10, "000000", None),
    ]
    values_col = [
        None, None, None, None,
        None,
        clean_res.original_rows,
        clean_res.final_rows,
        clean_res.duplicates_removed,
        None,
        None,
        clean_res.missing_before,
        clean_res.missing_after,
        None,
        None,
        clean_res.invalid_values,
        clean_res.corrected_categorical,
        clean_res.invalid_emails,
        None,
        None,
        ana.total_students,
        ana.avg_age,
        ana.avg_attendance,
        ana.avg_total,
        ana.max_total,
        ana.min_total,
        ana.low_attendance_count,
    ]

    for i, ((text, bold, size, color, bg), val) in enumerate(
        zip(rows, values_col), start=1
    ):
        c = ws.cell(row=i, column=1, value=text)
        c.font      = Font(bold=bold, size=size, color=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if bg:
            c.fill = PatternFill("solid", fgColor=bg)
        if val is not None:
            v_cell = ws.cell(row=i, column=2, value=val)
            v_cell.font      = Font(bold=True, size=10, color="1F4E79")
            v_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[i].height = 20

    ws.merge_cells("A1:B1")
    ws.merge_cells("A2:B2")
    ws.merge_cells("A3:B3")


def build_quality_report(ws, clean_res, ana):
    df = ana.df
    miss_before_series = pd.Series(dtype=int)

    # Rebuild approximate before-cleaning counts from the log
    before_map: dict[str, int] = {}
    for entry in clean_res.log:
        if entry["Step"] == "Missing Values":
            col = entry["Column"]
            before_map[col] = before_map.get(col, 0) + entry["Records_Affected"]

    rows_data = []
    for col in df.columns:
        miss_b    = before_map.get(col, 0)
        miss_a    = int(df[col].isna().sum())
        miss_pct  = round(miss_b / max(clean_res.original_rows, 1) * 100, 2)
        dup_count = int(df.duplicated(subset=[col]).sum())
        invalid   = sum(
            e["Records_Affected"] for e in clean_res.log
            if e["Column"] == col and "Flagged" in e["Action"]
        )
        action = " | ".join(
            e["Action"] for e in clean_res.log if e["Column"] == col
        ) or "No action required"
        rows_data.append({
            "Column":            col,
            "Data_Type":         str(df[col].dtype),
            "Missing_Before":    miss_b,
            "Missing_After":     miss_a,
            "Missing_%":         miss_pct,
            "Duplicate_Count":   dup_count,
            "Invalid_Values":    invalid,
            "Cleaning_Action":   action[:120],
        })

    _write_df_to_sheet(ws, pd.DataFrame(rows_data))


def build_department_analysis(ws, ana):
    dept = (
        ana.dept_counts
        .merge(ana.dept_avg_marks, on="Department")
        .merge(ana.dept_avg_attend, on="Department")
        .merge(ana.dept_max_marks, on="Department")
    )
    _write_df_to_sheet(ws, dept)


def build_grade_analysis(ws, ana):
    _write_df_to_sheet(ws, ana.grade_dist)


def build_cleaned_data(ws, ana):
    _write_df_to_sheet(ws, ana.df)


def build_cleaning_log(ws, clean_res):
    log_df = pd.DataFrame(clean_res.log)
    _write_df_to_sheet(ws, log_df)


def insert_image(ws, img_path: Path, anchor: str = "A1",
                 width: int = 500, height: int = 320):
    if img_path.exists():
        img = XLImage(str(img_path))
        img.width  = width
        img.height = height
        ws.add_image(img, anchor)


# ══════════════════════════════════════════════════════════════════════════════
#  MASTER REPORT BUILDER
# ══════════════════════════════════════════════════════════════════════════════

def generate_excel_report(clean_res, ana, chart_paths: dict[str, Path]):
    print("\n[Report] Building Excel workbook …")
    REPORT_FILE.parent.mkdir(parents=True, exist_ok=True)

    wb  = Workbook()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Sheet 1: Executive Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"
    build_executive_summary(ws1, clean_res, ana, now)

    # Sheet 2: Data Quality Report
    ws2 = wb.create_sheet("Data Quality Report")
    build_quality_report(ws2, clean_res, ana)

    # Sheet 3: Department Analysis
    ws3 = wb.create_sheet("Department Analysis")
    build_department_analysis(ws3, ana)

    # Sheet 4: Grade Analysis
    ws4 = wb.create_sheet("Grade Analysis")
    build_grade_analysis(ws4, ana)

    # Sheet 5: Cleaned Data
    ws5 = wb.create_sheet("Cleaned Data")
    build_cleaned_data(ws5, ana)

    # Sheet 6: Cleaning Log
    ws6 = wb.create_sheet("Data Cleaning Log")
    build_cleaning_log(ws6, clean_res)

    # Sheet 7: Charts (images embedded)
    ws7 = wb.create_sheet("Charts")
    ws7.sheet_view.showGridLines = False
    positions = {
        "dept_dist":    ("A1",  500, 340),
        "avg_marks":    ("J1",  600, 340),
        "grade_dist":   ("A21", 600, 340),
        "missing_vals": ("J21", 400, 340),
        "attendance":   ("A41", 500, 340),
        "dept_compare": ("J41", 500, 340),
    }
    for key, (anchor, w, h) in positions.items():
        if key in chart_paths:
            insert_image(ws7, chart_paths[key], anchor=anchor, width=w, height=h)

    wb.save(REPORT_FILE)
    print(f"    Excel report saved → {REPORT_FILE}")


# ══════════════════════════════════════════════════════════════════════════════
#  TEXT SUMMARY
# ══════════════════════════════════════════════════════════════════════════════

def generate_text_summary(clean_res, ana):
    now  = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = [
        "=" * 65,
        "  DATA CLEANING & REPORTING AUTOMATION – QUALITY SUMMARY",
        "=" * 65,
        f"  Generated on    : {now}",
        "",
        "  RECORD STATISTICS",
        f"  Original rows   : {clean_res.original_rows}",
        f"  Final rows      : {clean_res.final_rows}",
        f"  Duplicates removed: {clean_res.duplicates_removed}",
        "",
        "  MISSING VALUE STATISTICS",
        f"  Before cleaning : {clean_res.missing_before}",
        f"  After cleaning  : {clean_res.missing_after}",
        f"  Improvement     : {clean_res.missing_before - clean_res.missing_after}",
        "",
        "  DATA QUALITY",
        f"  Invalid values flagged      : {clean_res.invalid_values}",
        f"  Categorical values corrected: {clean_res.corrected_categorical}",
        f"  Invalid emails flagged      : {clean_res.invalid_emails}",
        "",
        "  ACADEMIC PERFORMANCE",
        f"  Total students     : {ana.total_students}",
        f"  Average age        : {ana.avg_age} years",
        f"  Average attendance : {ana.avg_attendance}%",
        f"  Average total marks: {ana.avg_total}",
        f"  Highest marks      : {ana.max_total}",
        f"  Lowest marks       : {ana.min_total}",
        f"  Students < 75% att : {ana.low_attendance_count}",
    ]
    if ana.correlation is not None:
        lines.append(f"  Attendance-Marks correlation: {ana.correlation}")

    lines += [
        "",
        "  DEPARTMENT BREAKDOWN",
    ]
    for _, row in ana.dept_counts.iterrows():
        lines.append(f"  {row['Department']:<6}: {row['Student_Count']} students")

    lines += [
        "",
        "  GRADE DISTRIBUTION",
    ]
    for _, row in ana.grade_dist.iterrows():
        lines.append(
            f"  Grade {row['Grade']:<3}: {row['Count']:>4} students  "
            f"({row['Percentage']:.1f}%)"
        )

    lines.append("=" * 65)
    text = "\n".join(lines)

    SUMMARY_TXT.parent.mkdir(parents=True, exist_ok=True)
    SUMMARY_TXT.write_text(text, encoding="utf-8")
    print(f"    Text summary saved → {SUMMARY_TXT}")
    return text


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════

def generate_reports(clean_res, ana):
    print("=" * 60)
    print("  Generating Reports")
    print("=" * 60)
    chart_paths = generate_all_charts(clean_res, ana)
    generate_excel_report(clean_res, ana, chart_paths)
    text = generate_text_summary(clean_res, ana)
    print("=" * 60)
    return chart_paths, text
